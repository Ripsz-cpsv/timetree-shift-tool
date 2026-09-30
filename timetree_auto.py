import time
from datetime import datetime, time as dtime
import os
import re
import sys
import traceback
import pandas as pd
from openpyxl import load_workbook
from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.common.keys import Keys
from selenium.webdriver.chrome.options import Options
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from selenium.common.exceptions import TimeoutException

# ==================== 【設定項目】 ====================
# 右側の "" の中に直接入力してください。
TIMETREE_EMAIL = os.environ.get("TIMETREE_EMAIL", "your_email@example.com")
TIMETREE_PASSWORD = os.environ.get("TIMETREE_PASSWORD", "sumple")
TARGET_CALENDAR_NAME = ""        # 抽出したいカレンダーの名前
TARGET_SHIFT_TITLE = ""                # 抽出したい予定のタイトル

OUTPUT_EXCEL_PATH = r"C:\Users\user\Downloads\時間自動連動_請求書・作業実績.xlsx"
TARGET_SHEET_NAME = "集計シート"

ZOOM = 0.5   # ブラウザの表示倍率（0.5 = 50%）。全予定が見えない場合は 0.4 などに下げる
# ======================================================

COLUMNS = ["日付", "予定タイトル", "開始時間", "終了時間", "労働時間(分)"]


def is_file_locked(filepath):
    """Excelファイルが開かれているか確認"""
    if not os.path.exists(filepath):
        return False
    try:
        with open(filepath, "r+b"):
            return False
    except OSError:
        return True


def extract_time(s):
    """文字列から時刻を取り出す（24時間表記・AM/PM表記どちらも対応）"""
    m = re.search(r"(\d{1,2}):(\d{2})\s*([AaPp][Mm])?", s or "")
    if not m:
        return None
    h, mi, ap = int(m.group(1)), int(m.group(2)), m.group(3)
    if ap:
        ap = ap.lower()
        if ap == "pm" and h < 12:
            h += 12
        elif ap == "am" and h == 12:
            h = 0
    if h > 23 or mi > 59:
        return None
    return dtime(h, mi)


def parse_datetime_str(raw_str):
    """
    TimeTreeモーダル内の日時文字列をパースする。
    戻り値: (日付文字列 or None, 時刻文字列 or None, datetime or None)
    """
    if not raw_str or raw_str == "未取得":
        return None, None, None

    clean = re.sub(r"^[A-Za-z]+,\s*", "", raw_str).strip()      # "Mon, " を除去
    clean = re.sub(r"[\(（][月火水木金土日][\)）]", "", clean)    # "(月)" を除去
    s = re.sub(r"\s+", " ", clean).strip()

    formats = [
        "%Y年%m月%d日 %H:%M",
        "%Y/%m/%d %H:%M",
        "%b %d %Y %I:%M%p",
        "%b %d %Y %I:%M %p",
        "%b %d %Y %H:%M",
        "%b %d, %Y %I:%M%p",
        "%b %d, %Y %I:%M %p",
        "%b %d, %Y %H:%M",
    ]
    for fmt in formats:
        try:
            dt = datetime.strptime(s, fmt)
            return dt.strftime("%Y/%m/%d"), dt.strftime("%H:%M"), dt
        except ValueError:
            continue

    # フォールバック: 日付と時刻を個別に抽出
    t = extract_time(clean)
    time_str = t.strftime("%H:%M") if t else None

    date_str, d_obj = None, None
    date_match = re.search(r"(\d{4})[/年](\d{1,2})[/月](\d{1,2})", clean)
    if date_match:
        try:
            d_obj = datetime(int(date_match.group(1)),
                             int(date_match.group(2)),
                             int(date_match.group(3)))
            date_str = d_obj.strftime("%Y/%m/%d")
        except ValueError:
            pass

    dt = datetime.combine(d_obj.date(), t) if (d_obj and t) else None
    return date_str, time_str, dt


def select_calendar(driver, wait, name):
    """左のリストから対象カレンダーを選択する"""
    # <li data-test-id="compact-calendar-list-item"><button><div role="img" aria-label="カレンダー名">
    xpath = (f"//li[@data-test-id='compact-calendar-list-item']"
             f"//button[.//div[@role='img' and @aria-label='{name}']]")
    try:
        el = wait.until(EC.presence_of_element_located((By.XPATH, xpath)))
        driver.execute_script("arguments[0].scrollIntoView({block: 'center'});", el)
        driver.execute_script("arguments[0].click();", el)
        time.sleep(2)
        return True
    except TimeoutException:
        names = [e.get_attribute("aria-label") for e in driver.find_elements(
            By.CSS_SELECTOR, "li[data-test-id='compact-calendar-list-item'] div[role='img']")]
        print(f"  ⚠️ カレンダー「{name}」が見つかりません。検出されたカレンダー: {names}")
        print("     現在表示中のカレンダーで続行します。")
        return False


def close_modal(driver, wait, start_elem):
    try:
        driver.find_element(By.CSS_SELECTOR, "button[aria-label='Close']").click()
    except Exception:
        webdriver.ActionChains(driver).send_keys(Keys.ESCAPE).perform()
    # 前のモーダルが消えるまで待つ（次の予定で古い値を拾わないため）
    try:
        WebDriverWait(driver, 5).until(EC.staleness_of(start_elem))
    except TimeoutException:
        pass


def write_to_excel(df_final):
    wb = load_workbook(OUTPUT_EXCEL_PATH)
    ws = wb[TARGET_SHEET_NAME] if TARGET_SHEET_NAME in wb.sheetnames else wb.create_sheet(TARGET_SHEET_NAME)

    # 前回のデータが残らないよう A〜E 列をクリア
    for row in ws.iter_rows(min_row=1, min_col=1, max_col=len(COLUMNS)):
        for cell in row:
            cell.value = None

    for c, col_name in enumerate(COLUMNS, start=1):
        ws.cell(row=1, column=c, value=col_name)
    for r, row in enumerate(df_final.itertuples(index=False), start=2):
        for c, val in enumerate(row, start=1):
            if isinstance(val, float) and pd.isna(val):
                val = None
            ws.cell(row=r, column=c, value=val)

    wb.save(OUTPUT_EXCEL_PATH)


print("🤖 自動化処理を開始します...")

if not TIMETREE_EMAIL or not TIMETREE_PASSWORD:
    print("❌ メールアドレス/パスワードが未設定です。環境変数か設定項目に入力してください。")
    input("\nエンターキーを押すと終了します...")
    sys.exit(1)

if is_file_locked(OUTPUT_EXCEL_PATH):
    print(f"⚠️ 警告: ファイル 「{OUTPUT_EXCEL_PATH}」 が開かれています。")
    print("👉 Excel等を閉じてから再度実行してください。")
    input("\nエンターキーを押すと終了します...")
    sys.exit(1)

chrome_options = Options()
chrome_options.add_argument("--window-size=1920,1080")
# ★ブラウザ全体を縮小表示（実際のビューポートが広がり、全予定が見える）
chrome_options.add_argument(f"--force-device-scale-factor={ZOOM}")

driver = webdriver.Chrome(options=chrome_options)
wait = WebDriverWait(driver, 15)

try:
    # --- ステップ 1: ログイン ---
    print("📌 [ステップ1/4] ログイン画面を開いています...")
    driver.get("https://timetreeapp.com/signin")

    email_field = wait.until(EC.presence_of_element_located((By.NAME, "email")))
    email_field.clear()
    email_field.send_keys(TIMETREE_EMAIL)

    pass_field = driver.find_element(By.NAME, "password")
    pass_field.clear()
    pass_field.send_keys(TIMETREE_PASSWORD)

    driver.find_element(By.XPATH, "//button[@type='submit']").click()

    wait.until(EC.presence_of_element_located((By.XPATH, "//main | //nav")))
    print("🚀 ログインが完了しました。")

    # --- ステップ 2: カレンダー選択 ---
    print(f"📌 [ステップ2/4] カレンダー「{TARGET_CALENDAR_NAME}」を選択中...")
    select_calendar(driver, wait, TARGET_CALENDAR_NAME)

    # --- ステップ 3: 各予定を開いて詳細を取得 ---
    print(f"📌 [ステップ3/4] 「{TARGET_SHIFT_TITLE}」（完全一致）の予定詳細を解析中...")

    xpath_target = f"//button[.//span[normalize-space(text())='{TARGET_SHIFT_TITLE}']]"

    wait.until(EC.presence_of_element_located((By.XPATH, xpath_target)))
    time.sleep(1)  # 描画完了待ち
    total_count = len(driver.find_elements(By.XPATH, xpath_target))
    print(f"🔍 画面上に {total_count} 件の「{TARGET_SHIFT_TITLE}」を発見しました。")

    records = []

    for idx in range(total_count):
        start_elem = None
        try:
            target_elems = driver.find_elements(By.XPATH, xpath_target)
            if idx >= len(target_elems):
                break
            elem = target_elems[idx]

            driver.execute_script("arguments[0].scrollIntoView({block: 'center'});", elem)
            time.sleep(0.2)
            driver.execute_script("arguments[0].click();", elem)

            start_elem = wait.until(
                EC.presence_of_element_located((By.CSS_SELECTOR, "[data-test-id='event-date-time-start']"))
            )
            raw_start = start_elem.text.replace("\n", " ").strip()

            try:
                end_elem = driver.find_element(By.CSS_SELECTOR, "[data-test-id='event-date-time-end']")
                raw_end = end_elem.text.replace("\n", " ").strip()
            except Exception:
                raw_end = "未取得"

            start_date, start_time, dt_start = parse_datetime_str(raw_start)
            end_date, end_time, dt_end = parse_datetime_str(raw_end)

            # 終了が「17:00」など時刻のみの場合は、開始日と組み合わせて補完
            if dt_start and not dt_end:
                t_end = extract_time(raw_end)
                if t_end:
                    dt_end = datetime.combine(dt_start.date(), t_end)
                    end_time = t_end.strftime("%H:%M")

            record_date = start_date or end_date or datetime.now().strftime("%Y/%m/%d")

            # 労働時間（分）
            labor_minutes = 0
            if dt_start and dt_end:
                labor_minutes = int((dt_end - dt_start).total_seconds() / 60)
                if labor_minutes < 0:
                    labor_minutes += 24 * 60  # 日跨ぎ補正

            records.append({
                "日付": record_date,
                "予定タイトル": TARGET_SHIFT_TITLE,
                "開始時間": start_time or raw_start,
                "終了時間": end_time or raw_end,
                "労働時間(分)": labor_minutes,
            })

            print(f"  └ [{idx + 1}/{total_count}] 解析成功: {record_date} | "
                  f"{TARGET_SHIFT_TITLE} | {start_time} ～ {end_time} ({labor_minutes}分)")

            close_modal(driver, wait, start_elem)
            time.sleep(0.3)

        except Exception as elem_err:
            print(f"  ⚠️ [{idx + 1}/{total_count}] 処理エラー: {elem_err}")
            try:
                webdriver.ActionChains(driver).send_keys(Keys.ESCAPE).perform()
            except Exception:
                pass
            time.sleep(0.5)

    print(f"\n📌 解析完了: 全 {len(records)} 件の処理が完了しました。")

    # --- ステップ 4: Excel出力（明細は分、最終行に合計h） ---
    print("📌 [ステップ4/4] Excelファイルへの書き込みを開始します...")
    if records:
        df = pd.DataFrame(records)[COLUMNS]
        df = df.sort_values("日付", kind="stable").reset_index(drop=True)

        total_minutes = int(df["労働時間(分)"].sum())
        total_hours = round(total_minutes / 60.0, 2)

        total_row = pd.DataFrame([{
            "日付": "合計",
            "予定タイトル": "",
            "開始時間": "",
            "終了時間": f"{total_hours}h",
            "労働時間(分)": None,
        }])
        df_final = pd.concat([df, total_row], ignore_index=True)[COLUMNS]

        if os.path.exists(OUTPUT_EXCEL_PATH):
            write_to_excel(df_final)
            print(f"✨ 成功！「{OUTPUT_EXCEL_PATH}」の「{TARGET_SHEET_NAME}」シートを更新しました。")
            print(f"📊 合計: {total_minutes}分 ({total_hours}時間)")
            print("\n【最終結果一覧】")
            print(df_final.to_string(index=False))
        else:
            print(f"❌ ファイルが存在しません: {OUTPUT_EXCEL_PATH}")
    else:
        print("❌ 該当する予定データが見つかりませんでした。")

except Exception as main_error:
    print("\n❌ エラー詳細:")
    print(f"エラー種別: {type(main_error).__name__}")
    print("--- スタックトレース（発生箇所） ---")
    traceback.print_exc()

finally:
    print("\n💡 確認のため画面を固定しています。エンターキーを押すと終了します...")
    input()
    driver.quit()